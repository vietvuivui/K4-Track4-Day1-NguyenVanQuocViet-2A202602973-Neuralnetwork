"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx.

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, không ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

COLUMNS = ["exp_id", "group", "description", "loss", "optimizer", "lr", "weight_decay", "batch", "epochs",
           "hidden", "dropout", "clip_norm", "precision", "init", "seed", "step0_loss", "best_val_loss",
           "best_epoch", "final_train_loss", "final_val_loss", "val_acc", "val_macro_f1", "time_per_epoch_s",
           "peak_mem_MB", "diverged", "eval_acc", "eval_macro_f1", "figure_file", "notes"]
FORMULA_COLUMNS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
MAX_ROWS = 60  # công thức của mẫu (Seeds, Summary) tham chiếu Experiments!2:61

# Giá trị theo danh sách chọn của mẫu
LOSS_NAMES = {"ce": "CE", "mse": "MSE"}
OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}


def _clean(obj):
    """Đổi tuple/numpy/NaN sang kiểu JSON được."""
    if isinstance(obj, dict):
        return {k: _clean(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_clean(v) for v in obj]
    if hasattr(obj, "item"):  # numpy / torch scalar
        obj = obj.item()
    if isinstance(obj, float) and not math.isfinite(obj):
        return None
    return obj


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    out = Path(results_dir)
    out.mkdir(parents=True, exist_ok=True)
    path = out / f"{result['cfg']['exp_id']}.json"
    payload = {k: _clean(result[k]) for k in ("cfg", "history", "summary")}
    for k in ("notes", "eval"):
        if k in result:
            payload[k] = _clean(result[k])
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    res = []
    for p in sorted(Path(results_dir).glob("*.json")):
        with open(p, encoding="utf-8") as f:
            res.append(json.load(f))
    return sorted(res, key=lambda r: r["cfg"]["exp_id"])


def _round(v, nd=4):
    return round(v, nd) if isinstance(v, float) else v


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng của bảng (khoá trùng tên cột).

    eval_scores: dict có "accuracy"/"macro_f1" (đúng như eval_result.json) — CHỈ cho baseline và cấu hình cuối.
    """
    cfg, s = result["cfg"], result["summary"]
    row = {
        "exp_id": cfg["exp_id"], "group": cfg.get("group", "other"), "description": cfg.get("description", ""),
        "loss": LOSS_NAMES.get(cfg["loss"], cfg["loss"]),
        "optimizer": OPT_NAMES.get(cfg["optimizer"], cfg["optimizer"]),
        "lr": cfg["lr"], "weight_decay": cfg.get("weight_decay", 0.0), "batch": cfg["batch"],
        "epochs": cfg["epochs"], "hidden": "-".join(str(h) for h in cfg["hidden"]),
        "dropout": cfg.get("dropout", 0.0),
        "clip_norm": "none" if cfg.get("clip_norm") is None else cfg["clip_norm"],
        "precision": cfg.get("precision", "fp32"), "init": cfg.get("init", "he"), "seed": cfg["seed"],
        "diverged": "Y" if s.get("diverged") else "N",
        "figure_file": f"figures/{cfg['exp_id']}.png",
    }
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss",
              "val_acc", "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[k] = _round(s.get(k), 2 if k in ("time_per_epoch_s", "peak_mem_MB") else 4)
    if eval_scores:
        row["eval_acc"] = _round(eval_scores.get("accuracy", eval_scores.get("eval_acc")))
        row["eval_macro_f1"] = _round(eval_scores.get("macro_f1", eval_scores.get("eval_macro_f1")))

    note_parts = [notes or result.get("notes", "")]
    if cfg.get("scheduler"):
        note_parts.append(f"scheduler={cfg['scheduler']}")
    if s.get("peak_mem_MB") is None:
        note_parts.append("peak_mem_MB trống: chạy trên CPU")
    row["notes"] = "; ".join(p for p in note_parts if p)
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), rồi lưu thành out_path.

    seed_ids     : exp_id các lần chạy baseline khác seed -> ghi vào Seeds!A2:A6 (tối đa 5).
    summary_notes: {group: "nhận xét"} -> ghi vào cột "nhận xét ngắn" của sheet Summary.
    Các cột công thức không bị ghi đè. Mở file bằng Excel/LibreOffice để công thức tính lại.
    """
    import openpyxl

    if len(rows) > MAX_ROWS:
        raise ValueError(f"mẫu chỉ có công thức cho {MAX_ROWS} dòng; hiện có {len(rows)} dòng")
    ids = [r["exp_id"] for r in rows]
    assert len(set(ids)) == len(ids), "exp_id bị trùng"

    wb = openpyxl.load_workbook(template_path)  # không data_only => giữ công thức
    ws = wb["Experiments"]
    header = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    missing = [c for c in COLUMNS if c not in header]
    assert not missing, f"mẫu thiếu cột {missing}"

    # xoá dữ liệu mẫu cũ (dòng baseline điền sẵn) ở các cột nhập liệu
    for r in range(2, MAX_ROWS + 2):
        for col in COLUMNS:
            ws.cell(row=r, column=header[col]).value = None

    for i, row in enumerate(rows):
        r = i + 2
        for col in COLUMNS:
            if col in FORMULA_COLUMNS:
                continue
            v = row.get(col)
            ws.cell(row=r, column=header[col]).value = v if v != "" else None

    if seed_ids:
        ws_s = wb["Seeds"]
        for i in range(5):
            ws_s.cell(row=2 + i, column=1).value = seed_ids[i] if i < len(seed_ids) else None

    if summary_notes:
        ws_m = wb["Summary"]
        hdr = {ws_m.cell(row=1, column=c).value: c for c in range(1, ws_m.max_column + 1)}
        note_col = next(c for name, c in hdr.items() if name and str(name).startswith("nhận xét"))
        for r in range(2, ws_m.max_row + 1):
            g = ws_m.cell(row=r, column=1).value
            if g in summary_notes:
                ws_m.cell(row=r, column=note_col).value = summary_notes[g]

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
